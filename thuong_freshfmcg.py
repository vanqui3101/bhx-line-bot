"""
thuong_freshfmcg.py - Báo cáo THƯỞNG FRESH/FMCG theo Base cố định (MỚI, tách
riêng khỏi lệnh TD/THƯỞNG cũ).

Luồng hoạt động:
1. Đọc TRỰC TIẾP file Excel "doanh thu chi tiết" (file SKU-level đã dùng
   chung cho MTKM/TD — mỗi dòng 1 sản phẩm/1 ngày, có cột "Ngày xuất",
   "Ngành hàng", "Thành tiền phải thu khách hàng (chưa VAT)"), KHÔNG phụ
   thuộc excel_reader.py để tránh đụng code cũ.
2. Tách Fresh/FMCG theo cột "Ngành hàng":
   FRESH = Rau Củ Các Loại / Trái Cây Các Loại / Thịt gia cầm gia súc các
   loại / Thủy Hải Sản Các Loại. Còn lại = FMCG.
3. Cộng dồn theo NGÀY, lưu vào SQLite (bảng freshfmcg_daily, khóa=ngày) —
   gửi file mới đè đúng ngày đó, các ngày khác giữ nguyên (cộng dồn qua
   nhiều lần gửi, không cần gửi lại từ đầu tháng).
4. Khi gõ lệnh "TTFF" / "tiến độ thưởng": tính lũy kế đầu tháng->hiện tại,
   % so Base cố định (Fresh 373.000.000đ / FMCG 1.056.000.000đ), thưởng
   lũy kế + thưởng dự kiến cuối tháng (kiểu DTDK: TB/ngày đã qua x số ngày
   cả tháng), theo ĐÚNG công thức đã chốt với anh Quí:
     - DT < 80% Base           -> thưởng = 0
     - 80% <= DT <= 100% Base  -> (DT - 80%Base) x tỉ lệ mốc 1
     - DT > 100% Base          -> 20%Base x tỉ lệ mốc 1  +  (DT - Base) x tỉ lệ mốc 2
       (tỉ lệ mốc 2 = gấp đôi tỉ lệ mốc 1, cộng dồn với phần mốc 1)
   Fresh: tỉ lệ mốc 1 = 1,4%  | mốc 2 = 2,8%
   FMCG : tỉ lệ mốc 1 = 0,7%  | mốc 2 = 1,4%
"""
import calendar
from datetime import datetime
import storage

FRESH_NGANH_HANG = {
    "Rau Củ Các Loại",
    "Trái Cây Các Loại",
    "Thịt gia cầm gia súc các loại",
    "Thủy Hải Sản Các Loại",
}

BASE_FRESH = 373_000_000
BASE_FMCG = 1_056_000_000
TY_LE_FRESH_MOC1 = 0.014
TY_LE_FRESH_MOC2 = 0.028
TY_LE_FMCG_MOC1 = 0.007
TY_LE_FMCG_MOC2 = 0.014


def _chuan_hoa_ngay(v):
    """Chuẩn hoá 1 ô ngày (datetime object do openpyxl tự parse, hoặc
    string "dd/mm/yyyy") về "YYYY-MM-DD". Trả về None nếu không đọc được."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        v = v.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(v, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def doc_freshfmcg_theo_ngay(file_path):
    """Đọc file 'doanh thu chi tiết' (.xlsx), trả về dict:
    { "YYYY-MM-DD": {"fresh": tổng_tiền, "fmcg": tổng_tiền}, ... }
    Tự tìm cột theo TÊN CỘT (không phụ thuộc thứ tự cột trong file)."""
    from openpyxl import load_workbook
    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    header = [str(h).strip() if h is not None else "" for h in header]
    idx = {name: i for i, name in enumerate(header)}
    col_ngay = idx.get("Ngày xuất")
    col_nganh = idx.get("Ngành hàng")
    col_tien = idx.get("Thành tiền phải thu khách hàng (chưa VAT)")
    if col_ngay is None or col_nganh is None or col_tien is None:
        raise ValueError(
            "File thiếu cột bắt buộc — cần đủ 3 cột: 'Ngày xuất', 'Ngành hàng', "
            "'Thành tiền phải thu khách hàng (chưa VAT)'. Anh kiểm tra lại đúng "
            "file \"doanh thu chi tiết\" giúp em."
        )
    max_col = max(col_ngay, col_nganh, col_tien)
    per_day = {}
    for row in rows_iter:
        if row is None or len(row) <= max_col:
            continue
        ngay_raw = row[col_ngay]
        nganh = row[col_nganh]
        tien = row[col_tien]
        if ngay_raw is None or tien is None:
            continue
        ngay_str = _chuan_hoa_ngay(ngay_raw)
        if ngay_str is None:
            continue
        try:
            tien_f = float(tien)
        except (TypeError, ValueError):
            continue
        la_fresh = (str(nganh).strip() if nganh is not None else "") in FRESH_NGANH_HANG
        bucket = per_day.setdefault(ngay_str, {"fresh": 0.0, "fmcg": 0.0})
        if la_fresh:
            bucket["fresh"] += tien_f
        else:
            bucket["fmcg"] += tien_f
    wb.close()
    return per_day


def extract_and_save(file_path):
    """Đọc file rồi lưu (upsert theo ngày) vào SQLite. Trả về số ngày đã đọc được."""
    per_day = doc_freshfmcg_theo_ngay(file_path)
    if not per_day:
        return 0
    rows = [
        {"ngay": ngay, "fresh_dt": v["fresh"], "fmcg_dt": v["fmcg"]}
        for ngay, v in per_day.items()
    ]
    storage.save_freshfmcg_daily(rows)
    return len(rows)


def tinh_thuong(dt, base, ty_le_moc1, ty_le_moc2):
    """Công thức thưởng đã chốt — xem docstring đầu file."""
    if dt <= 0 or base <= 0:
        return 0.0
    nguong_80 = 0.8 * base
    if dt <= nguong_80:
        return 0.0
    if dt <= base:
        return (dt - nguong_80) * ty_le_moc1
    thuong_moc1 = (base - nguong_80) * ty_le_moc1
    thuong_moc2 = (dt - base) * ty_le_moc2
    return thuong_moc1 + thuong_moc2


def _khoi_nhom(dt_luy_ke, dt_hom_nay, base, ty_le_moc1, ty_le_moc2, ngay_hien_tai, so_ngay_ca_thang):
    du_kien_cuoi_thang = (dt_luy_ke / ngay_hien_tai * so_ngay_ca_thang) if ngay_hien_tai else 0.0
    return {
        "hom_nay": dt_hom_nay,
        "luy_ke": dt_luy_ke,
        "base": base,
        "pct_base": (dt_luy_ke / base * 100) if base else 0.0,
        "thuong_luy_ke": tinh_thuong(dt_luy_ke, base, ty_le_moc1, ty_le_moc2),
        "du_kien_cuoi_thang": du_kien_cuoi_thang,
        "thuong_du_kien": tinh_thuong(du_kien_cuoi_thang, base, ty_le_moc1, ty_le_moc2),
    }


def build_report_payload(ten_st):
    """Trả về payload đầy đủ để dựng thẻ báo cáo, hoặc None nếu chưa có data.
    Lấy đúng THÁNG của ngày có data mới nhất (không cố định 'tháng hiện tại'
    theo giờ hệ thống), để vẫn đúng nếu anh gửi data trễ."""
    rows = storage.get_freshfmcg_all()
    if not rows:
        return None
    ngay_moi_nhat = max(r["ngay"] for r in rows)
    nam, thang = int(ngay_moi_nhat[:4]), int(ngay_moi_nhat[5:7])
    prefix = f"{nam:04d}-{thang:02d}"
    rows_thang = sorted([r for r in rows if r["ngay"].startswith(prefix)], key=lambda r: r["ngay"])
    if not rows_thang:
        return None
    ngay_hien_tai_trong_thang = int(ngay_moi_nhat[8:10])
    so_ngay_ca_thang = calendar.monthrange(nam, thang)[1]
    hom_nay_row = rows_thang[-1]
    fresh_luy_ke = sum(r["fresh_dt"] for r in rows_thang)
    fmcg_luy_ke = sum(r["fmcg_dt"] for r in rows_thang)
    return {
        "ten_st": ten_st,
        "ngay_moi_nhat": ngay_moi_nhat,
        "ngay_hien_tai_trong_thang": ngay_hien_tai_trong_thang,
        "so_ngay_ca_thang": so_ngay_ca_thang,
        "fresh": _khoi_nhom(
            fresh_luy_ke, hom_nay_row["fresh_dt"], BASE_FRESH,
            TY_LE_FRESH_MOC1, TY_LE_FRESH_MOC2, ngay_hien_tai_trong_thang, so_ngay_ca_thang,
        ),
        "fmcg": _khoi_nhom(
            fmcg_luy_ke, hom_nay_row["fmcg_dt"], BASE_FMCG,
            TY_LE_FMCG_MOC1, TY_LE_FMCG_MOC2, ngay_hien_tai_trong_thang, so_ngay_ca_thang,
        ),
    }


# ---------------------------------------------------------------------------
# Dựng thẻ Flex Message (tự dựng JSON, KHÔNG phụ thuộc flex_builder.py để
# tránh đụng code cũ) — đúng layout đã duyệt: header vàng gold, 3 KPI trên
# cùng, 2 khối Fresh (xanh lá)/FMCG (xanh dương) có thanh % tiến độ.
# ---------------------------------------------------------------------------
def _money(n):
    return f"{n:,.0f}".replace(",", ".") + " đ"


def _kpi_box(number_text, label_text):
    return {
        "type": "box", "layout": "vertical", "backgroundColor": "#F7F7F7",
        "cornerRadius": "8px", "paddingAll": "10px", "flex": 1,
        "contents": [
            {"type": "text", "text": number_text, "size": "sm", "weight": "bold", "align": "center", "wrap": True},
            {"type": "text", "text": label_text, "size": "xxs", "color": "#777777", "align": "center", "margin": "4px", "wrap": True},
        ],
    }


def _row_line(label, value_text, value_color="#1A1A1A"):
    return {
        "type": "box", "layout": "horizontal", "margin": "3px",
        "contents": [
            {"type": "text", "text": label, "size": "xs", "color": "#1A1A1A", "flex": 3},
            {"type": "text", "text": value_text, "size": "xs", "weight": "bold", "color": value_color, "align": "end", "flex": 2, "wrap": True},
        ],
    }


def _thanh_tien_do(pct):
    pct_clamped = max(0.0, min(100.0, pct))
    return {
        "type": "box", "layout": "horizontal", "backgroundColor": "#FFFFFF",
        "height": "8px", "cornerRadius": "6px", "margin": "6px",
        "contents": [
            {"type": "box", "layout": "vertical", "backgroundColor": "#1E8246",
             "flex": int(round(pct_clamped)), "contents": []},
            {"type": "box", "layout": "vertical", "backgroundColor": "#FFFFFF",
             "flex": int(round(100 - pct_clamped)), "contents": []},
        ],
    }


def _khoi_nhom_box(ten_nhom, khoi, mau_nen, mau_thanh):
    bar = _thanh_tien_do(khoi["pct_base"])
    bar["contents"][0]["backgroundColor"] = mau_thanh
    thuong_luy_ke_text = _money(khoi["thuong_luy_ke"]) if khoi["thuong_luy_ke"] > 0 else "0 đ (chưa qua 80%)"
    return {
        "type": "box", "layout": "vertical", "backgroundColor": mau_nen,
        "cornerRadius": "8px", "paddingAll": "12px", "margin": "10px",
        "contents": [
            {"type": "text", "text": ten_nhom, "weight": "bold", "size": "sm", "color": "#1A1A1A"},
            _row_line("Lũy kế đầu tháng", _money(khoi["luy_ke"])),
            _row_line("Base tháng (cố định)", _money(khoi["base"])),
            bar,
            _row_line("% so Base", f"{khoi['pct_base']:.1f}%"),
            _row_line("Thưởng lũy kế hiện tại", thuong_luy_ke_text, mau_thanh),
            _row_line("DT dự kiến cuối tháng", _money(khoi["du_kien_cuoi_thang"])),
            _row_line("Thưởng DỰ KIẾN cuối tháng", _money(khoi["thuong_du_kien"]), mau_thanh),
        ],
    }


def build_freshfmcg_flex_message(payload):
    """Trả về bubble (dict) — anh Quí ghép với FlexMessage/FlexContainer như
    các báo cáo khác trong app.py."""
    ngay_disp = datetime.strptime(payload["ngay_moi_nhat"], "%Y-%m-%d").strftime("%d/%m/%Y")
    return {
        "type": "bubble",
        "size": "mega",
        "header": {
            "type": "box", "layout": "vertical", "backgroundColor": "#F6D365", "paddingAll": "14px",
            "contents": [
                {"type": "text", "text": "💰 TIẾN ĐỘ THƯỞNG FRESH & FMCG", "weight": "bold", "size": "md", "color": "#1A1A1A", "wrap": True},
            ],
        },
        "body": {
            "type": "box", "layout": "vertical", "paddingAll": "14px",
            "contents": [
                {"type": "text", "text": f"{payload['ten_st']}  |  Cập nhật đến {ngay_disp}", "size": "xxs", "color": "#777777", "wrap": True},
                {
                    "type": "box", "layout": "horizontal", "margin": "10px", "spacing": "sm",
                    "contents": [
                        _kpi_box(_money(payload["fresh"]["hom_nay"] + payload["fmcg"]["hom_nay"]), "DT hôm nay"),
                        _kpi_box(_money(payload["fresh"]["luy_ke"] + payload["fmcg"]["luy_ke"]), "DT lũy kế đầu tháng"),
                        _kpi_box(f"{payload['ngay_hien_tai_trong_thang']}/{payload['so_ngay_ca_thang']}", "ngày đã qua"),
                    ],
                },
                _khoi_nhom_box("FRESH", payload["fresh"], "#E4F3E9", "#1E8246"),
                _khoi_nhom_box("FMCG", payload["fmcg"], "#E4EDFA", "#1A66CC"),
                {"type": "text", "text": "* Base cố định không đổi theo tháng. Thưởng: <80% Base = 0đ, 80-100% x tỉ lệ mốc 1, vượt 100% phần vượt x tỉ lệ mốc 2 (cộng dồn).",
                 "size": "xxs", "color": "#999999", "margin": "10px", "wrap": True},
            ],
        },
    }
